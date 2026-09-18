import copy
from pathlib import Path

from docx import Document
from docx.table import _Row
from docx.text.paragraph import Paragraph

SOURCE_COMPLETE = Path(r"C:\winword\ccc\CarePortal Note Code mydoc complete Latest.docx")
SOURCE_SHRINK = Path(r"C:\winword\ccc\CarePortal Note Code mydoc shrink Latest.docx")
OUTPUT_COMPLETE = Path(r"C:\winword\aaa\AutoISF CarePortal Note Code Full Reference mydoc Sep 19 26 DRAFT updates.docx")
OUTPUT_SHRINK = Path(r"C:\winword\aaa\AutoISF CarePortal Note Code Abbreviated Reference mydoc Sep 19 26 DRAFT updates.docx")

# (anchor code already in the doc, new code, full description) -- insert new code immediately
# AFTER anchor in the existing alphabetical Table 1 / paragraph flow. LStOff anchors off LStOn
# itself so the pair lands together in the order they're inserted.
FULL_ENTRIES = [
    ("DuraR", "EDSR", "EarlyDawnSlowRise (04:30-08:00): sustained slow-rise entry criteria "
                       "(delta/shortAvgDelta/longAvgDelta all 0-0.15mmol) -> SMBdel +0.5, a 30-minute switch "
                       "to the TierC-slot profile, and a 30-minute TT."),
    ("LocOn", "LoReb", "recentLowReboundGuardFiredThisCycle (DetermineBasalAutoISF.kt): halves the current "
                        "microBolus after a recent low; no throttle, 60-minute lookback."),
    ("LOGF3", "LStOn", "ApsAutoIsfUseLiveStepsOnVirtual toggled ON: VirtualPump mirrors the loop phone's "
                        "live step counts via NS instead of reading its own (non-existent) local pedometer."),
    ("LStOn", "LStOff", "ApsAutoIsfUseLiveStepsOnVirtual toggled OFF: VirtualPump reverts to its own "
                         "local step source."),
    ("MJrec", "MJs2", "MjStateMj2TT: List1 direct action manually forcing the MJ automation state to MJ2."),
    ("MJs3", "MJsAc", "MjStateActiveTT: List1 direct action manually forcing the MJ automation state to "
                       "MJ active."),
    ("MoreM", "MoreMJ2", "Hypoalarm-path variant of MoreMJ: advances MJ state to MJ2 specifically when a "
                          "hypo alarm is active (more cautious than the plain MoreMJ progression)."),
]

SHRINK_ENTRIES = [
    ("DuraR", "EDSR", "EarlyDawnSlowRise: 04:30-08:00 slow-rise"),
    ("LocOn", "LoReb", "recentLowReboundGuard: halves microBolus"),
    ("LOGF3", "LStOn", "Live steps on VirtualPump: mirrors live"),
    ("LStOn", "LStOff", "Live steps on VirtualPump: local source"),
    ("MJrec", "MJs2", "MjStateMj2TT: manually forces MJ2"),
    ("MJs3", "MJsAc", "MjStateActiveTT: manually forces MJ active"),
    ("MoreM", "MoreMJ2", "MoreMJ2 (hypoalarm): advances to MJ2"),
]


def find_row_by_code(table, code):
    for row in table.rows:
        if row.cells[0].text.strip() == code:
            return row
    raise ValueError(f"anchor code not found in table: {code!r}")


def set_cell_text(cell, text):
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p0 = cell.paragraphs[0]
    runs = p0.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        p0.add_run(text)


def insert_row_after(table, anchor_code, code, description):
    anchor_row = find_row_by_code(table, anchor_code)
    new_tr = copy.deepcopy(anchor_row._tr)
    anchor_row._tr.addnext(new_tr)
    new_row = _Row(new_tr, table)
    for cell, text in zip(new_row.cells[:3], [code, f"\u2022 {code}", description]):
        set_cell_text(cell, text)


def splice_complete():
    doc = Document(SOURCE_COMPLETE)
    table1 = doc.tables[1]
    for anchor, code, desc in FULL_ENTRIES:
        insert_row_after(table1, anchor, code, desc)
    doc.add_paragraph(
        "Table 0 above (the two-stream quick-lookup index, general codes left / MJ-family codes right) was "
        "NOT updated in this pass -- only Table 1 (the single authoritative alphabetical list) was spliced. "
        "Flag if Table 0 should also be kept in sync."
    )
    doc.save(OUTPUT_COMPLETE)


def find_para_by_code(doc, code):
    for p in doc.paragraphs:
        first_field = p.text.split("\t", 1)[0].strip()
        if first_field == code:
            return p
    raise ValueError(f"anchor code not found in paragraphs: {code!r}")


def insert_para_after(doc, anchor_code, code, description):
    anchor_p = find_para_by_code(doc, anchor_code)
    new_p_elem = copy.deepcopy(anchor_p._p)
    anchor_p._p.addnext(new_p_elem)
    new_p = Paragraph(new_p_elem, anchor_p._parent)
    text = f"{code}\t{code}\t{description}"
    runs = new_p.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        new_p.add_run(text)


def splice_shrink():
    doc = Document(SOURCE_SHRINK)
    for anchor, code, desc in SHRINK_ENTRIES:
        insert_para_after(doc, anchor, code, desc)
    doc.save(OUTPUT_SHRINK)


splice_complete()
splice_shrink()
print(OUTPUT_COMPLETE)
print(OUTPUT_SHRINK)
