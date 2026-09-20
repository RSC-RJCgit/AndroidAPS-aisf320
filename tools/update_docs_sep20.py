"""20 Sep 2026 doc update: everything changed since the 19 Sep 2331 files (manual v29 / CarePortal Full v3 / registry v10).

New dated/versioned files only.  Manual v29 -> v30 (raw XML), CarePortal Full v3 -> v4 and Abbreviated v2 -> v3 (python-docx).
Changes covered: carb-split rounding-leftover fix + deferred last-zero rule (BolusWizard); FastRise size-tier and LoReb test
toggles (Settings + List 2 5.226 / 5.228, notes FRtOn/FRtOff/LRbOn/LRbOff, RT reason lines); RawMiss tag on Virtual AIV rows;
HiBrk brakes switched off and back on (no net change)."""
import copy
import glob
import re
import zipfile
from datetime import datetime
from html import escape
from pathlib import Path

from docx import Document
from docx.table import _Row
from docx.text.paragraph import Paragraph

AAA = Path(r"C:\winword\aaa")
stamp = f"{datetime.now():%H%M}"

# ------------------------------------------------------------------ manual
SRC = Path(sorted(glob.glob(str(AAA / "AutoISF Operations FULL Manual revised Saturday, September 19th, 2026 mydoc v29 * fully sorted careportal tables.docx")))[-1])
OUT = AAA / f"AutoISF Operations FULL Manual revised Sunday, September 20th, 2026 mydoc v30 {stamp} fully sorted careportal tables.docx"
x = zipfile.ZipFile(SRC).read("word/document.xml").decode("utf-8")
PARA_RE = re.compile(r"<w:p[ >].*?</w:p>", re.S)


def ptext(p):
    return "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", p))


def para_xml(lead, body, indent=True):
    ppr = '<w:pPr><w:ind w:left="720"/><w:jc w:val="both"/></w:pPr>' if indent else '<w:pPr><w:jc w:val="both"/></w:pPr>'
    return (f"<w:p>{ppr}<w:r><w:rPr><w:b/></w:rPr><w:t xml:space=\"preserve\">{escape(lead, quote=False)} </w:t></w:r>"
            f"<w:r><w:t xml:space=\"preserve\">{escape(body, quote=False)}</w:t></w:r></w:p>")


def insert_after_prefix(prefix, new_xml, label, which="all"):
    """Insert new_xml after every paragraph whose text starts with prefix (which='all'), asserting at least one hit."""
    global x
    hits = [m for m in PARA_RE.finditer(x) if ptext(m.group(0)).startswith(prefix)]
    assert hits, (label, "no hit")
    if which == "unique":
        assert len(hits) == 1, (label, len(hits))
    for m in reversed(hits):
        x = x[: m.end()] + new_xml + x[m.end():]
    print("ok:", label, f"({len(hits)} place(s))")


SPLIT = (
    "Carb split (BolusWizard.scheduleReducedPartsSplitBolus): (a) rounding-leftover fix. The residual is the calculated total minus "
    "the max bolus, and 5.2 minus 5.0 is 0.20000000000000018 in floating point, so after the 0.2 U part was delivered a remainder of "
    "about 1.7e-16 U still counted as pending and kept scheduling a check every interval, each one writing a 0.00 U bolus, calc row and "
    "S0.00 note until the 60-minute deadline. The remainder is now rounded to 0.001 U, and anything under half a pump step ends the "
    "split quietly (log line only). (b) Deferred zeros: a part skipped because IOB rose (calculated 0 U) no longer writes its 0 U bolus, "
    "calc row and S0.00 note at once. Only the last skipped part is written, once, and only if the split then ends without delivering "
    "(60-minute deadline, stop, superseding entry, profile switch below 100%, pump suspended, superbolus, or three consecutive unsafe BG "
    "checks); a later delivered part discards the remembered skip. Delivered parts keep their own bolus record and calc row. The "
    "protein/fat FPU series is unchanged: only its last dose writes a zero marker."
)
insert_after_prefix("Confirmation and visible history. Pressing OK performs one final recalculation", para_xml("Updated 20 Sep 2026.", SPLIT), "carb split", "unique")

RAWMISS = (
    "On VirtualPump only (not the Client app), the AIV Notes column, in both the on-screen table and the CSV/TXT export, gets the tag "
    "RawMiss when no raw value above 10 sits within 3 minutes of the row. The raw-dependent automations (BMild, bg3, Tier 3's entry "
    "gate, HP1, the HiBrk UKF and HP gates) read the raw value that the BG source piggybacks in GlucoseValue.noise, so when Virtual's BG "
    "source stops delivering it (for example an xDrip broadcast with no Extras.Raw: rawBGL 0.0, RawUKF5 \u201c--\u201d) they silently "
    "cannot fire; the tag makes that visible. Not yet built or device-verified."
)
insert_after_prefix("Updated 4 Sep 2026. The on-screen AIV table dropped UKF3", para_xml("Updated 20 Sep 2026.", RAWMISS), "rawmiss", "unique")

TOGGLES = (
    "Two test toggles. (a) FastRise SMB tiers: ApsAutoIsfFastRiseEnabled, default ON (behaviour unchanged). Off skips only the three "
    "Libre FAST RISE size-tier conditions in determine_basal (the main x0.2-x0.9 cascade, the early-rise x0.7 tier and the higher-BG "
    "x0.75 tier); the sensor-glitch guards, the early-AM and twilight limits, the late FastRise taper and the 30-minute cumulative SMB "
    "cap stay active. While off, the RT reason carries \u201cFastRise tiers OFF (test)\u201d. (b) Post-low rebound guard (LoReb): "
    "ApsAutoIsfLowReboundGuardEnabled, default OFF (it was hard-coded off on 20 Sep and became this setting the same day). While off it "
    "never fires, so there is no LoReb note, LastLoRebAppliedAt is not refreshed (only StuckRisingSlowly reads it, as a recent-event "
    "exclusion) and the RT reason carries \u201cLoReb OFF (test)\u201d. Both are per-phone switches in the AutoISF settings and List 2 "
    "rows (5.226 / 5.228)."
)
insert_after_prefix("One further wrinkle worth flagging for anyone reading exported FastRise values", para_xml("Updated 20 Sep 2026.", TOGGLES), "toggles", "unique")

LIST2 = (
    "Two new rows: 5.226 \u201cFastRise SMB tiers on/off (test)\u201d and 5.228 \u201cPost-low rebound guard LoReb on/off (test)\u201d. "
    "Each shows its current value, flips ApsAutoIsfFastRiseEnabled / ApsAutoIsfLowReboundGuardEnabled immediately on the loop phone "
    "(SMS plus CarePortal note FRtOn / FRtOff and LRbOn / LRbOff), and on AAPSClient relays the TT code to the loop phone, where the "
    "invoke-time FastRiseToggleTT / LowReboundGuardToggleTT blocks apply it. Two existing rows were missing from this catalogue: 5.206 "
    "\u201cLive steps on VirtualPump on/off\u201d (VirtualPump only) and 5.208 \u201cADB wireless: attempt Shizuku start\u201d (Virtual only)."
)
insert_after_prefix("Updated 6 Sep 2026 evening. List 2 Newest NNN looks under", para_xml("Updated 20 Sep 2026.", LIST2), "list2 (body + appendix E)")

CHANGELOG = (
    "Carb split rounding-leftover loop fixed and carb-split zero markers deferred to the last skip only (BolusWizard); FastRise size-tier "
    "and LoReb test toggles (Settings and List 2 5.226 / 5.228, notes FRtOn/FRtOff/LRbOn/LRbOff); RawMiss tag on Virtual AIV rows. The "
    "HiBrk day/night brakes were switched off and back on the same day, so the 19 Sep rework stands unchanged."
)
insert_after_prefix("2026-09-19. HiBrk / HiBrkDay / HiBrkDayMid reworked", para_xml("2026-09-20.", CHANGELOG, indent=False), "changelog", "unique")

with zipfile.ZipFile(SRC) as zin, zipfile.ZipFile(OUT, "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/document.xml":
            data = x.encode("utf-8")
        zout.writestr(item, data, compress_type=item.compress_type)
print(OUT)

# ------------------------------------------------------------------ CarePortal Full v3 -> v4
SRC_F = Path(sorted(glob.glob(str(AAA / "AutoISF CarePortal Note Code Full Reference mydoc Sep 19 26 DRAFT updates v3 *.docx")))[-1])
OUT_F = AAA / f"AutoISF CarePortal Note Code Full Reference mydoc Sep 20 26 DRAFT updates v4 {stamp}.docx"
NEW_FULL = [
    ("FRtOn", "FRtOn", "FastRise SMB size tiers switched ON (List 2 5.226 or the Settings switch). Default state."),
    ("FRtOf", "FRtOff", "FastRise SMB size tiers switched OFF (test): the Libre FAST RISE tiers are skipped; glitch guards, late taper and cumulative SMB cap stay active."),
    ("LRbOn", "LRbOn", "Post-low rebound guard (LoReb) switched ON (List 2 5.228 or the Settings switch)."),
    ("LRbOf", "LRbOff", "Post-low rebound guard (LoReb) switched OFF (test). Default state; the guard never fires while off."),
    ("RawMi", "RawMiss (AIV tag)", "AIV Notes-column tag on VirtualPump rows only (not a CarePortal note): no raw value above 10 within 3 minutes, so BMild, bg3, Tier 3 entry, HP1 and the HiBrk gates cannot fire."),
]
doc = Document(SRC_F)
table = doc.tables[0]


def set_cell_text(cell, text):
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for extra in p0.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        p0.add_run(text)


for code, display, desc in NEW_FULL:
    rows = table.rows
    anchor = next((r for r in rows if r.cells[0].text.strip().lower() > code.lower()), None)
    template = anchor if anchor is not None else rows[-1]
    new_tr = copy.deepcopy(template._tr)
    if anchor is not None:
        anchor._tr.addprevious(new_tr)
    else:
        template._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    for cell, text in zip(row.cells[:3], [code, f"\u2022 {display}", desc]):
        set_cell_text(cell, text)
doc.save(OUT_F)
print(OUT_F)

# ------------------------------------------------------------------ CarePortal Abbreviated v2 -> v3
SRC_A = Path(sorted(glob.glob(str(AAA / "AutoISF CarePortal Note Code Abbreviated Reference mydoc Sep 19 26 DRAFT updates v2 *.docx")))[-1])
OUT_A = AAA / f"AutoISF CarePortal Note Code Abbreviated Reference mydoc Sep 20 26 DRAFT updates v3 {stamp}.docx"
NEW_ABBR = [
    ("FRtOn", "FRtOn", "FastRise tiers ON (List 2 5.226)"),
    ("FRtOf", "FRtOff", "FastRise tiers OFF (test)"),
    ("LRbOn", "LRbOn", "LoReb guard ON (List 2 5.228)"),
    ("LRbOf", "LRbOff", "LoReb guard OFF (test)"),
    ("RawMi", "RawMiss", "AIV tag, Virtual: no raw data"),
]
doc = Document(SRC_A)
for code, display, desc in NEW_ABBR:
    paras = doc.paragraphs
    anchor = next((p for p in paras[1:] if p.text.split("\t", 1)[0].strip().lower() > code.lower() and "\t" in p.text), None)
    assert anchor is not None, code
    elem = copy.deepcopy(anchor._p)
    anchor._p.addprevious(elem)
    para = Paragraph(elem, anchor._parent)
    text = f"{code}\t{display}\t{desc}"
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        para.add_run(text)
doc.save(OUT_A)
print(OUT_A)
