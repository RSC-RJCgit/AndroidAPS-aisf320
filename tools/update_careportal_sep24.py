"""24 Sep 2026 CarePortal update: Full v4 -> v5 and Abbreviated v3 -> v4 (new dated/versioned files only).

New rows: HiBrkTwilight, HiBrkTwilightCut, T3uOn/T3uOff (List 2 5.230).  Refreshed rows: EDSR (24 Sep criteria), HnAM (window)."""
import copy
import glob
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.table import _Row
from docx.text.paragraph import Paragraph

AAA = Path(r"C:\winword\aaa")
stamp = f"{datetime.now():%H%M}"


def set_cell_text(cell, text):
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for extra in p0.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        p0.add_run(text)


# ------------------------------------------------------------------ Full v4 -> v5
SRC_F = Path(sorted(glob.glob(str(AAA / "AutoISF CarePortal Note Code Full Reference mydoc Sep 20 26 DRAFT updates v4 *.docx")))[-1])
assert "DELL17" not in SRC_F.name and "Latest" not in SRC_F.name
OUT_F = AAA / f"AutoISF CarePortal Note Code Full Reference mydoc Sep 24 26 DRAFT updates v5 {stamp}.docx"
NEW_FULL = [
    ("HBTwi", "HiBrkTwilight", "Dawn high-glucose brake (04:00-07:00, BG 6.5-9.0), TT-only variant: TT 4.2 mmol for 2 min, own 15-minute throttle; no SMB delivery boost, pp weight or profile action. Entry needs no active TT, positive raw UKF deltas, dura the dominant ISF factor, IOB change < 0.5 U in 5 min and HP1 >= 6.2 (lowered from 6.5 on 24 Sep 2026; not yet device-verified)."),
    ("HBTCt", "HiBrkTwilightCut", "HiBrkTwilight's own TT cancelled early: a delta stopped being positive/in range or IOB rose by 1.0 U or more in 5 minutes."),
    ("T3uOn", "T3uOn / T3uOff", "Tier 3 UAM boost \u201cunrestricted\u201d compare toggle switched on or off (List 2 5.230 or the Settings switch). On skips the 9.5 per-SMB cap, the 20-minute / SMBdel > 0.50 cuts and the percent-IOB gate, for A/B comparison only (default off)."),
]
UPDATED_FULL = {
    "EDSR": "EarlyDawnSlowRise (04:30-08:00): sustained slow rise (delta, shortAvgDelta and longAvgDelta all above 0 and below 0.15 mmol) with BG above 6.5 (was 7.0 until 24 Sep 2026) and no active TT -> TT 4.4 mmol for 5 minutes (was 30), own 5-minute throttle so it re-arms only while the slow rise persists. No SMB delivery or profile action. The 24 Sep change is not yet device-verified.",
    "HnAM": "\"HighNight00AM\": sets the Profile state to HnAM for the re-enabled night-high / OffHighProf tightened-exit pair. Window 02:00-07:00 (since 24 Sep 2026).",
}
doc = Document(SRC_F)
table = doc.tables[0]
for r in table.rows:
    key = r.cells[0].text.strip()
    if key in UPDATED_FULL:
        set_cell_text(r.cells[2], UPDATED_FULL[key])
# The Full table is sorted by the note name in column 2 (HiBrkDayMid < HiBrkTwilight; T3FastCarb before T3uOn), not by the
# 5-character code in column 1, so each new row goes directly after a named existing/just-added row.
AFTER = {"HiBrkTwilight": "HiBrkDayMid", "HiBrkTwilightCut": "HiBrkTwilight", "T3uOn / T3uOff": "T3FastCarb"}
for code, display, desc in NEW_FULL:
    anchor = next(r for r in table.rows if r.cells[1].text.strip() == f"• {AFTER[display]}")
    new_tr = copy.deepcopy(anchor._tr)
    anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    for cell, text in zip(row.cells[:3], [code, f"• {display}", desc]):
        set_cell_text(cell, text)
doc.save(OUT_F)
print(OUT_F)

# ------------------------------------------------------------------ Abbreviated v3 -> v4
SRC_A = Path(sorted(glob.glob(str(AAA / "AutoISF CarePortal Note Code Abbreviated Reference mydoc Sep 20 26 DRAFT updates v3 *.docx")))[-1])
assert "DELL17" not in SRC_A.name and "Latest" not in SRC_A.name
OUT_A = AAA / f"AutoISF CarePortal Note Code Abbreviated Reference mydoc Sep 24 26 DRAFT updates v4 {stamp}.docx"
NEW_ABBR = [
    ("HBTwi", "HiBrkTwilight", "Dawn TT 4.2 for 2 min (04:00-07:00)"),
    ("HBTCt", "HiBrkTwilightCut", "HiBrkTwilight TT cut short"),
    ("T3uOn", "T3uOn / T3uOff", "Tier 3 unrestricted toggle (List 2 5.230)"),
]
UPDATED_ABBR = {
    "EDSR": "EarlyDawnSlowRise: 04:30-08:00 slow rise, BG > 6.5, TT 4.4 for 5 min",
    "HnAM": "HighNight00AM, 02:00-07:00",
}
doc = Document(SRC_A)
for p in doc.paragraphs:
    if "\t" in p.text:
        key = p.text.split("\t", 1)[0].strip()
        if key in UPDATED_ABBR:
            disp = p.text.split("\t")[1]
            text = f"{key}\t{disp}\t{UPDATED_ABBR[key]}"
            p.runs[0].text = text
            for extra in p.runs[1:]:
                extra._r.getparent().remove(extra._r)
for code, display, desc in NEW_ABBR:
    paras = doc.paragraphs
    anchor = next((p for p in paras[1:] if p.text.split("\t", 1)[0].strip().lower() > code.lower() and "\t" in p.text), None)
    assert anchor is not None, code
    elem = copy.deepcopy(anchor._p)
    anchor._p.addprevious(elem)
    para = Paragraph(elem, anchor._parent)
    text = f"{code}\t{display}\t{desc}"
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        para.add_run(text)
doc.save(OUT_A)
print(OUT_A)
